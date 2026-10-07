#!/usr/bin/env python3
"""Regenerates the fixtures for scripts/test-phase45-*.ts from a real Word authoring library (python-docx)."""
import os, struct, zlib
from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
os.chdir(HERE)
KEY_MARK = "ZZ-ANSWER-KEY-7491"

def tiny_png(path):
    def chunk(t, d): return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xffffffff)
    raw = b"".join(b"\x00" + b"\xff\x80\x00" * 4 for _ in range(4))
    open(path, "wb").write(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 4, 4, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))

def lesson(with_image=False):
    d = Document()
    d.add_heading("Subtracting Fractions with Unlike Denominators", level=1)
    d.add_heading("Learning Objectives", level=2)
    for o in ["Find a common denominator for two fractions.", "Subtract fractions with unlike denominators and simplify the result."]:
        d.add_paragraph(o, style="List Bullet")
    d.add_heading("Introduction", level=2)
    d.add_paragraph("Fractions can only be subtracted when their pieces are the same size. If the denominators differ, we first rename both fractions so they share a denominator.")
    d.add_paragraph("The least common multiple of the two denominators is the most convenient common denominator.")
    if with_image:
        tiny_png("pic.png"); d.add_picture("pic.png"); os.remove("pic.png")
    d.add_heading("Worked Example 1", level=2)
    d.add_paragraph("Subtract 3/4 - 1/6. The least common multiple of 4 and 6 is 12, so 3/4 = 9/12 and 1/6 = 2/12. Then 9/12 - 2/12 = 7/12.")
    d.add_heading("Practice", level=2)
    for q in ["1. Compute 5/6 - 1/4.", "2. Compute 7/8 - 2/3.", "3. A recipe uses 3/4 cup of sugar and 1/3 cup is already added. How much more is needed?"]:
        d.add_paragraph(q)
    d.add_heading("Key Points", level=2)
    d.add_paragraph("Rename first, subtract the numerators second, simplify last.")
    d.add_heading("Answer Key", level=2)
    d.add_paragraph(f"{KEY_MARK}: 1) 7/12  2) 5/24  3) 5/12 cup")
    return d

lesson().save("lesson-structured.docx")
lesson(True).save("lesson-with-image.docx")

# many sections (limit test)
d = Document(); d.add_heading("A very long lesson", level=1)
for i in range(1, 46):
    d.add_heading(f"Section {i}", level=2); d.add_paragraph(f"This is the text of section {i}; it explains one small idea in a complete sentence.")
d.save("too-many-sections.docx")

# one giant section (split test)
d = Document(); d.add_heading("A long explanation", level=1); d.add_heading("Explanation", level=2)
for i in range(1, 120):
    d.add_paragraph(f"Paragraph {i}. " + "The idea builds on the previous one and adds a new detail. " * 6)
d.save("long-section.docx")

# only an answer key
d = Document(); d.add_heading("Answer Key", level=1); d.add_paragraph(f"{KEY_MARK}: 1) 4  2) 9  3) 12")
d.save("only-answer-key.docx")

open("lesson-plain.txt", "w").write("""Percent of a Number

LEARNING OBJECTIVES
- Find a percent of a number.
- Use percents in everyday situations.

INTRODUCTION
A percent is a ratio out of 100. To find 20% of 80, write 20% as 0.20 and multiply: 0.20 x 80 = 16.

Example 1: Finding a discount
A shirt costs 400 pesos and is 25% off. The discount is 0.25 x 400 = 100 pesos, so the sale price is 300 pesos.

Practice
1. What is 10% of 250?
2. What is 35% of 60?

SUMMARY
Percent means per hundred. Convert to a decimal, then multiply.

ANSWER KEY
ZZ-ANSWER-KEY-7491: 1) 25  2) 21
""")
print("fixtures written")
